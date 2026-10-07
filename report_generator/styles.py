"""
styles.py: Định nghĩa các kiểu dáng chuẩn học thuật cho tài liệu Word bằng python-docx.
Quy chuẩn:
- Font chữ: Times New Roman
- Cỡ chữ văn bản thường: 13pt
- Giãn dòng: 1.35 lines
- Spacing Before / After: 3pt / 3pt
- Căn lề trang: Trái 3.0 cm, Phải 2.0 cm, Trên 2.0 cm, Dưới 2.0 cm
- Heading 1: 16pt In đậm
- Heading 2: 14pt In đậm
- Heading 3: 13pt In đậm nghiêng
"""

import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

COLOR_PRIMARY = RGBColor(0, 0, 0)
COLOR_MUTED = RGBColor(80, 80, 80)
COLOR_HEADER_BG = "EAECEE"
COLOR_BORDER = "BDC3C7"


def setup_document_styles(doc):
    """Cấu hình kích thước trang và các style mặc định cho tài liệu."""
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    # Cấu hình Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = COLOR_PRIMARY
    style_normal.paragraph_format.line_spacing = 1.35
    style_normal.paragraph_format.space_before = Pt(3)
    style_normal.paragraph_format.space_after = Pt(3)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    return section


def add_paragraph(doc, text="", bold_prefix=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=3, space_after=3):
    """Thêm một đoạn văn bản chuẩn học thuật, hỗ trợ in đậm tiền tố và căn lề."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)

    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Times New Roman'
        r_prefix.font.size = Pt(13)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = COLOR_PRIMARY

    if text:
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(13)
        r_text.font.italic = italic
        r_text.font.color.rgb = COLOR_PRIMARY

    return p


def add_heading_1(doc, text):
    """Thêm tiêu đề cấp 1: 16pt, In đậm."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return p


def add_heading_2(doc, text):
    """Thêm tiêu đề cấp 2: 14pt, In đậm."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return p


def add_heading_3(doc, text):
    """Thêm tiêu đề cấp 3: 13pt, In đậm nghiêng."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = COLOR_PRIMARY
    return p


def add_caption(doc, text, is_table=True):
    """Thêm nhãn chú thích bảng hoặc hình: 12pt, In nghiêng đậm."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if is_table else 3)
    p.paragraph_format.space_after = Pt(4 if is_table else 8)
    p.paragraph_format.keep_with_next = True if is_table else False
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = COLOR_PRIMARY
    return p


def add_formula_block(doc, formula_text, formula_id=None):
    """Thêm khối công thức toán học căn giữa kèm số thứ tự công thức căn phải."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(formula_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.italic = True

    if formula_id:
        p_num = doc.add_paragraph()
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_num.paragraph_format.space_before = Pt(0)
        p_num.paragraph_format.space_after = Pt(4)
        r_num = p_num.add_run(formula_id)
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(12)
        r_num.font.italic = True
    return p


def set_cell_background(cell, fill_hex):
    """Thiết lập màu nền cho ô trong bảng."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    """Thiết lập đường viền thanh mảnh, trang nhã cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)


def add_styled_table(doc, headers, rows_data, col_widths=None, alignment=None):
    """Tạo bảng dữ liệu chuẩn học thuật với tiêu đề nổi bật và định dạng dòng xen kẽ."""
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], COLOR_HEADER_BG)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY

    # Format Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F9FAFA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            if bg_color != "FFFFFF":
                set_cell_background(row_cells[c_idx], bg_color)
            p = row_cells[c_idx].paragraphs[0]
            if alignment and c_idx < len(alignment):
                p.alignment = alignment[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                run.font.color.rgb = COLOR_PRIMARY

    # Thiết lập độ rộng cột nếu có
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Cm(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table
