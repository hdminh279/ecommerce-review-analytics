"""
main.py: Tập tin chính kích hoạt sinh Báo Cáo Toàn Diện và Hoàn Chỉnh (.docx).
Tạo tài liệu, gọi các hàm xây dựng từng chương và lưu file Bao_Cao_Hoan_Chinh.docx.
"""

import os
import sys
import docx
from report_generator.styles import setup_document_styles
from report_generator.front_matter import build_cover_page, build_preface_and_outlines
from report_generator.chapter1 import build_chapter_1
from report_generator.chapter2 import build_chapter_2
from report_generator.chapter3 import build_chapter_3
from report_generator.chapter4 import build_chapter_4_and_references


def generate_full_report(output_filename="Bao_Cao_Hoan_Chinh.docx"):
    print(f"[*] Bắt đầu quá trình khởi tạo tài liệu {output_filename}...")
    doc = docx.Document()

    # 1. Cấu hình kiểu dáng và kích thước trang chuẩn A4
    print("    -> Cấu hình margins (Trái 3cm, Phải 2cm, Trên 2cm, Dưới 2cm) và style chữ...")
    setup_document_styles(doc)

    # 2. Xây dựng trang bìa
    print("    -> Đang tạo Trang bìa chính quy...")
    build_cover_page(doc)

    # 3. Lời nói đầu, Tính cấp thiết ban đầu, Mục tiêu, Danh mục viết tắt
    print("    -> Đang tạo Lời mở đầu, Mục tiêu và Danh mục viết tắt...")
    build_preface_and_outlines(doc)

    # 4. Chương 1: Tổng quan về đề tài và cơ sở dữ liệu
    print("    -> Đang biên soạn CHƯƠNG 1 (Tổng quan về đề tài và cơ sở dữ liệu)...")
    build_chapter_1(doc)

    # 5. Chương 2: Cơ sở lý thuyết và công nghệ nền tảng
    print("    -> Đang biên soạn CHƯƠNG 2 (Cơ sở lý thuyết và công nghệ nền tảng)...")
    build_chapter_2(doc)

    # 6. Chương 3: Thiết kế hệ thống và thực nghiệm
    print("    -> Đang biên soạn CHƯƠNG 3 (Thiết kế hệ thống và thực nghiệm)...")
    build_chapter_3(doc)

    # 7. Chương 4: Kết luận và hướng phát triển + Tài liệu tham khảo
    print("    -> Đang biên soạn CHƯƠNG 4 & TÀI LIỆU THAM KHẢO...")
    build_chapter_4_and_references(doc)

    # 8. Lưu tệp
    output_path = os.path.abspath(output_filename)
    doc.save(output_path)
    print(f"[+] Xuất bản thành công tệp: {output_path}")

    # 9. Thống kê số lượng từ và đoạn văn
    total_words = 0
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)

    for p in doc.paragraphs:
        words = p.text.strip().split()
        total_words += len(words)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    words = p.text.strip().split()
                    total_words += len(words)

    print("\n" + "=" * 60)
    print("THỐNG KÊ KỸ THUẬT VĂN BẢN ĐÃ TẠO:")
    print(f" - Tổng số đoạn văn (paragraphs): {total_paragraphs}")
    print(f" - Tổng số bảng biểu (tables):      {total_tables}")
    print(f" - Tổng dung lượng từ (word count): {total_words:,} từ tiếng Việt")
    print("=" * 60)

    return output_path, total_words


if __name__ == "__main__":
    generate_full_report()
