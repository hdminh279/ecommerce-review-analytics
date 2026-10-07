"""
front_matter.py: Xây dựng trang bìa chính quy, lời mở đầu, mục tiêu nghiên cứu,
Mục lục tổng thể, Danh mục từ viết tắt, Danh mục bảng biểu và Danh mục hình vẽ.
Tuân thủ quy tắc: Không lạm dụng dấu ngoặc đơn, văn phong học thuật tiếng Việt trang trọng.
"""

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_generator.styles import (
    add_paragraph, add_heading_1, add_heading_2, add_styled_table
)


def build_cover_page(doc):
    """Tạo trang bìa chuẩn theo mẫu Đề cương Đồ án Tốt nghiệp Đại học Thủy Lợi."""
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(0)
    p_univ.paragraph_format.space_after = Pt(2)
    r1 = p_univ.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC THỦY LỢI\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(36)
    r_line = p_line.add_run("----------------***----------------")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(12)

    p_title_sub = doc.add_paragraph()
    p_title_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_title_sub.add_run("BÁO CÁO TOÀN DIỆN VÀ HOÀN CHỈNH\nĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(40)
    r_title = p_title.add_run("ĐỀ TÀI:\nXÂY DỰNG HỆ THỐNG ĐÁNH GIÁ TÂM LÝ NGƯỜI DÙNG VÀ PHÁT HIỆN XU HƯỚNG BẤT THƯỜNG TRÊN DỮ LIỆU THƯƠNG MẠI ĐIỆN TỬ")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(17)
    r_title.font.bold = True

    p_major = doc.add_paragraph()
    p_major.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_major.paragraph_format.space_after = Pt(45)
    r_major = p_major.add_run("Chuyên ngành: Công nghệ Thông tin\nMã ngành: 7480201")
    r_major.font.name = "Times New Roman"
    r_major.font.size = Pt(13)
    r_major.font.italic = True

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Cm(3.5)
    p_info.paragraph_format.space_after = Pt(80)
    p_info.paragraph_format.line_spacing = 1.35
    
    r_info = p_info.add_run(
        "Giảng viên hướng dẫn:  ThS. Trần Thị Ngà\n"
        "Bộ môn:               Khoa Công nghệ Thông tin\n\n"
        "Sinh viên thực hiện:   Hồ Đức Minh\n"
        "Mã số sinh viên:       2151173760\n"
        "Lớp chuyên ngành:      63TH2\n"
    )
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(13)
    r_info.font.bold = True

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_after = Pt(0)
    r_foot = p_foot.add_run("Hà Nội, Năm 2026")
    r_foot.font.name = "Times New Roman"
    r_foot.font.size = Pt(13)
    r_foot.font.bold = True

    doc.add_page_break()


def build_preface_and_outlines(doc):
    """Xây dựng phần tính cấp thiết, mục tiêu ban đầu theo đúng đề cương chuẩn."""
    add_heading_1(doc, "1. TÍNH CẤP THIẾT CỦA ĐỀ TÀI")
    
    add_paragraph(doc, 
        "Trong những năm gần đây, sự phát triển bùng nổ của cuộc Cách mạng Công nghiệp lần thứ tư cùng làn sóng chuyển đổi số mạnh mẽ trong lĩnh vực bán lẻ đã làm thay đổi căn bản hành vi mua sắm của người tiêu dùng trên toàn cầu. Hàng triệu lượt giao dịch thương mại điện tử diễn ra mỗi ngày sản sinh ra một khối lượng dữ liệu khổng lồ các bản đánh giá, nhận xét và phản hồi trực tuyến. Nguồn dữ liệu văn bản phi cấu trúc này chứa đựng những thông tin vô giá phản ánh chân thực tâm lý, thái độ, mức độ thỏa mãn cũng như trải nghiệm sử dụng thực tế của khách hàng đối với từng dòng sản phẩm. Tuy nhiên, tốc độ gia tăng nhanh chóng cùng quy mô dữ liệu vượt ngưỡng hàng triệu bản ghi mỗi ngày đã đặt ra thách thức vượt quá khả năng tổng hợp, rà soát và đánh giá thủ công của con người [1]."
    )

    add_paragraph(doc, 
        "Trong thực tiễn vận hành doanh nghiệp bán lẻ và các nhà phân phối lớn, các bộ phận kiểm soát chất lượng sản phẩm cùng bộ phận chăm sóc khách hàng thường xuyên rơi vào thế bị động, phản ứng rất chậm trễ trước các đợt bùng phát lỗi hàng loạt. Các vấn đề nghiêm trọng như lỗi linh kiện gia công, hỏng hóc trong khâu vận chuyển, sản phẩm kích ứng da hoặc các sai lệch về mô tả bao bì thường chỉ được phát hiện khi người dùng đồng loạt khiếu nại gay gắt hoặc tiến hành hoàn trả hàng với tỷ lệ đột biến. Nguyên nhân cốt lõi bắt nguồn từ việc doanh nghiệp thiếu hụt các công cụ kỹ thuật có khả năng giám sát tâm lý khách hàng tự động, liên tục và phát hiện sớm các dấu hiệu bất thường tiềm ẩn trong chuỗi thời gian. Khi sự cố được phát hiện theo phương thức thụ động truyền thống, thiệt hại về uy tín thương hiệu, niềm tin người tiêu dùng và chi phí tài chính đã chạm mức đặc biệt nghiêm trọng [2]."
    )

    add_paragraph(doc, 
        "Bên cạnh yêu cầu về nghiệp vụ kinh doanh, bài toán xử lý dữ liệu lớn phi cấu trúc còn đòi hỏi sự kết hợp chặt chẽ giữa các nền tảng kỹ thuật dữ liệu phân tán hiện đại và các thuật toán học máy xử lý ngôn ngữ tự nhiên, song hành cùng phương pháp kiểm soát quá trình bằng thống kê chuyên sâu. Xuất phát từ những đòi hỏi cấp bách của thực tiễn khoa học và kinh doanh, đề tài “Xây dựng hệ thống đánh giá tâm lý người dùng và phát hiện xu hướng bất thường trên dữ liệu thương mại điện tử” được triển khai nhằm xây dựng một giải pháp kỹ thuật tổng thể, tự động hóa xuyên suốt từ khâu tiếp nhận dữ liệu thô, phân loại sắc thái cảm xúc, phát hiện sớm các sản phẩm có nguy cơ lỗi cao cho đến việc cung cấp bảng thông tin trực quan hỗ trợ nhà quản trị ra quyết định kịp thời [3]."
    )

    add_heading_1(doc, "2. MỤC TIÊU NGHIÊN CỨU ĐỀ TÀI")

    add_paragraph(doc,
        "Đề tài hướng tới việc xây dựng một hệ thống đường ống dữ liệu phân tán toàn diện và tự động hóa cao độ, phục vụ trọn vẹn quy trình tiếp nhận dữ liệu đánh giá sản phẩm Amazon, làm sạch dữ liệu lớn, phân loại cảm xúc văn bản theo ba trạng thái chuẩn tắc, mô hình hóa xu hướng biến động theo chuỗi thời gian, và cảnh báo chủ động các sản phẩm phát sinh đột biến lỗi cần can thiệp kỹ thuật. Hệ thống được kiến trúc hóa theo mô hình Medallion Lakehouse ba tầng tiêu chuẩn gồm tầng Đồng, tầng Bạc và tầng Vàng, kết hợp công cụ mô hình hóa dữ liệu dbt và hệ quản trị cơ sở dữ liệu phân tích nhúng DuckDB, xuất dữ liệu phục vụ bảng điều khiển tương tác trên nền tảng Streamlit [4]."
    )

    doc.add_page_break()

    # MỤC LỤC CHI TIẾT
    add_heading_1(doc, "MỤC LỤC TỔNG THỂ")
    headers_toc = ["Nội dung chuyên đề / Mục", "Cấp độ"]
    rows_toc = [
        ["1. TÍNH CẤP THIẾT CỦA ĐỀ TÀI", "Phần mở đầu"],
        ["2. MỤC TIÊU NGHIÊN CỨU ĐỀ TÀI", "Phần mở đầu"],
        ["DANH MỤC CÁC TỪ VIẾT TẮT", "Mục tham chiếu"],
        ["DANH MỤC CÁC BẢNG BIỂU", "Mục tham chiếu"],
        ["DANH MỤC CÁC HÌNH VẼ VÀ SƠ ĐỒ", "Mục tham chiếu"],
        ["CHƯƠNG 1. TỔNG QUAN VỀ ĐỀ TÀI VÀ CƠ SỞ DỮ LIỆU", "Chương 1"],
        ["   1.1. Tính cấp thiết của đề tài", "Tiểu mục 1.1"],
        ["   1.2. Mục tiêu nghiên cứu (Mục tiêu tổng quát và kỹ thuật)", "Tiểu mục 1.2"],
        ["   1.3. Đối tượng và phạm vi nghiên cứu", "Tiểu mục 1.3"],
        ["      1.3.1. Cấu trúc tập dữ liệu đánh giá người dùng", "Tiểu mục 1.3.1"],
        ["      1.3.2. Cấu trúc tập dữ liệu mô tả sản phẩm", "Tiểu mục 1.3.2"],
        ["      1.3.3. Thách thức cố hữu của dữ liệu thực tế", "Tiểu mục 1.3.3"],
        ["   1.4. Kết luận chương", "Tiểu mục 1.4"],
        ["CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ NỀN TẢNG", "Chương 2"],
        ["   2.1. Xử lý ngôn ngữ tự nhiên và phân tích cảm xúc", "Tiểu mục 2.1"],
        ["      2.1.1. Tổng quan về bài toán phân tích cảm xúc", "Tiểu mục 2.1.1"],
        ["      2.1.2. Kỹ thuật tiền xử lý văn bản", "Tiểu mục 2.1.2"],
        ["      2.1.3. Kỹ thuật biểu diễn văn bản: TF-IDF và N-grams", "Tiểu mục 2.1.3"],
        ["      2.1.4. Thuật toán phân loại học máy: Linear SVM và Hồi quy Logistic", "Tiểu mục 2.1.4"],
        ["      2.1.5. Kỹ thuật giải quyết mất cân bằng nhãn", "Tiểu mục 2.1.5"],
        ["   2.2. Kiến trúc Big Data và hệ thống tính toán phân tán", "Tiểu mục 2.2"],
        ["      2.2.1. Tổng quan về Apache Spark và cơ chế tính toán trong bộ nhớ", "Tiểu mục 2.2.1"],
        ["      2.2.2. PySpark MLlib và Tối ưu hóa Catalyst", "Tiểu mục 2.2.2"],
        ["      2.2.3. Kiến trúc Medallion Lakehouse", "Tiểu mục 2.2.3"],
        ["      2.2.4. Định dạng lưu trữ Parquet và Phân vùng Partitioning", "Tiểu mục 2.2.4"],
        ["   2.3. Hạ tầng lưu trữ và điều phối", "Tiểu mục 2.3"],
        ["      2.3.1. Hệ thống lưu trữ đối tượng MinIO", "Tiểu mục 2.3.1"],
        ["      2.3.2. Công cụ mô hình hóa dữ liệu dbt và In-Process OLAP DuckDB", "Tiểu mục 2.3.2"],
        ["      2.3.3. Lý thuyết Kiểm soát chất lượng thống kê: Z-Score và Cửa sổ trượt", "Tiểu mục 2.3.3"],
        ["      2.3.4. Đóng gói hệ thống với Docker và Trực quan hóa Streamlit", "Tiểu mục 2.3.4"],
        ["   2.4. Kết luận chương", "Tiểu mục 2.4"],
        ["CHƯƠNG 3. THIẾT KẾ HỆ THỐNG VÀ THỰC NGHIỆM", "Chương 3"],
        ["   3.1. Phân tích dữ liệu thô và thiết kế kiến trúc tổng thể", "Tiểu mục 3.1"],
        ["      3.1.1. Phân tích tập dữ liệu thô Amazon 2023", "Tiểu mục 3.1.1"],
        ["      3.1.2. Thiết kế kiến trúc hệ thống tổng thể", "Tiểu mục 3.1.2"],
        ["   3.2. Xây dựng luồng tiền xử lý dữ liệu", "Tiểu mục 3.2"],
        ["      3.2.1. Tiếp nhận và nạp dữ liệu vào tầng Đồng", "Tiểu mục 3.2.1"],
        ["      3.2.2. Làm sạch và chuẩn hóa dữ liệu tại tầng Bạc bằng PySpark", "Tiểu mục 3.2.2"],
        ["   3.3. Huấn luyện, đánh giá và lựa chọn mô hình Sentiment Engine", "Tiểu mục 3.3"],
        ["      3.3.1. Chiến lược phân chia dữ liệu và phòng chống rò rỉ", "Tiểu mục 3.3.1"],
        ["      3.3.2. Thực nghiệm đối chuẩn toàn diện các phương pháp Sentiment Engine", "Tiểu mục 3.3.2"],
        ["      3.3.3. Lựa chọn Mô hình Vô địch và Đánh giá trên Hold-Out Test Set", "Tiểu mục 3.3.3"],
        ["      3.3.4. Lưu trữ mô hình và Suy luận hàng loạt vào tầng Vàng", "Tiểu mục 3.3.4"],
        ["   3.4. Phân tích xu hướng và phát hiện sản phẩm tiêu cực bằng dbt-DuckDB", "Tiểu mục 3.4"],
        ["      3.4.1. Thiết kế các tầng Staging, Intermediate và Data Marts", "Tiểu mục 3.4.1"],
        ["      3.4.2. Tính toán xu hướng biến động theo tháng của Category", "Tiểu mục 3.4.2"],
        ["      3.4.3. Xếp hạng Top sản phẩm phát sinh khiếu nại nhiều nhất", "Tiểu mục 3.4.3"],
        ["      3.4.4. Động cơ cảnh báo đột biến lỗi sản phẩm theo Z-Score SPC", "Tiểu mục 3.4.4"],
        ["      3.4.5. Quyết định lược bỏ khai phá khía cạnh để tối ưu hóa kiến trúc", "Tiểu mục 3.4.5"],
        ["   3.5. Xây dựng giao diện trực quan hóa và cảnh báo Streamlit", "Tiểu mục 3.5"],
        ["      3.5.1. Thiết kế kiến trúc Lock-Free và bộ đệm In-Memory", "Tiểu mục 3.5.1"],
        ["      3.5.2. Chi tiết các chức năng nghiệp vụ trên Bảng điều khiển", "Tiểu mục 3.5.2"],
        ["   3.6. Đánh giá hiệu năng và kết quả thực nghiệm", "Tiểu mục 3.6"],
        ["      3.6.1. Đánh giá chất lượng mô hình học máy và phân tích lỗi", "Tiểu mục 3.6.1"],
        ["      3.6.2. Đánh giá hiệu năng kỹ thuật dữ liệu và độ trễ hệ thống", "Tiểu mục 3.6.2"],
        ["   3.7. Kết luận chương", "Tiểu mục 3.7"],
        ["CHƯƠNG 4. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "Chương 4"],
        ["   4.1. Kết quả đạt được của đề tài", "Tiểu mục 4.1"],
        ["   4.2. Những hạn chế còn tồn tại", "Tiểu mục 4.2"],
        ["   4.3. Hướng nghiên cứu và phát triển tiếp theo", "Tiểu mục 4.3"],
        ["TÀI LIỆU THAM KHẢO", "Tham khảo"]
    ]
    add_styled_table(doc, headers_toc, rows_toc, col_widths=[12.0, 4.0])

    doc.add_page_break()

    # DANH MỤC BẢNG BIỂU
    add_heading_2(doc, "DANH MỤC CÁC BẢNG BIỂU")
    headers_lot = ["Ký hiệu bảng", "Tên bảng dữ liệu biểu diễn", "Chương"]
    rows_lot = [
        ["Bảng 1.1", "Cấu trúc các trường thông tin trong tập dữ liệu Amazon Reviews 2023", "Chương 1"],
        ["Bảng 1.2", "Cấu trúc các trường thông tin trong tập dữ liệu Product Metadata", "Chương 1"],
        ["Bảng 2.1", "So sánh cơ chế tính toán giữa Hadoop MapReduce và Apache Spark", "Chương 2"],
        ["Bảng 3.1", "Thống kê quy mô phân chia tập dữ liệu huấn luyện, kiểm định và kiểm thử", "Chương 3"],
        ["Bảng 3.2", "Bảng đối chuẩn kết quả thực nghiệm các phương pháp Sentiment Engine trên tập Validation", "Chương 3"],
        ["Bảng 3.3", "Kết quả đánh giá Mô hình Vô địch trên tập Hold-Out Test độc lập (431.487 mẫu)", "Chương 3"],
        ["Bảng 3.4", "Ma trận nhầm lẫn (Confusion Matrix) của Mô hình Vô địch trên tập Hold-Out Test", "Chương 3"],
        ["Bảng 3.5", "Cấu trúc các bảng dữ liệu tầng Data Marts xây dựng bằng dbt và DuckDB", "Chương 3"],
        ["Bảng 3.6", "Thống kê hiệu năng thực thi của toàn bộ đường ống dữ liệu", "Chương 3"]
    ]
    add_styled_table(doc, headers_lot, rows_lot, col_widths=[3.0, 10.5, 2.5])

    # DANH MỤC HÌNH VẼ
    add_heading_2(doc, "DANH MỤC CÁC HÌNH VẼ VÀ SƠ ĐỒ")
    headers_lof = ["Ký hiệu hình", "Tên hình vẽ và sơ đồ kiến trúc", "Chương"]
    rows_lof = [
        ["Sơ đồ 2.1", "Kiến trúc luân chuyển dữ liệu ba tầng Medallion Lakehouse (Bronze -> Silver -> Gold)", "Chương 2"],
        ["Sơ đồ 2.2", "Quy trình tối ưu hóa truy vấn thông minh bốn giai đoạn của Catalyst Optimizer", "Chương 2"],
        ["Sơ đồ 3.1", "Kiến trúc tổng thể hệ thống End-to-End từ tiếp nhận dữ liệu thô đến Dashboard", "Chương 3"],
        ["Sơ đồ 3.2", "Cơ chế phát hiện đột biến tỷ lệ khiếu nại Z-score SPC trên cửa sổ trượt ba tháng quá khứ", "Chương 3"],
        ["Sơ đồ 3.3", "Mô hình kết nối chỉ đọc không khóa tệp và bộ đệm RAM của Streamlit với DuckDB", "Chương 3"]
    ]
    add_styled_table(doc, headers_lof, rows_lof, col_widths=[3.0, 10.5, 2.5])

    # DANH MỤC TỪ VIẾT TẮT
    add_heading_2(doc, "DANH MỤC CÁC TỪ VIẾT TẮT")
    headers_abbr = ["Từ viết tắt", "Thuật ngữ tiếng Anh nguyên bản", "Ý nghĩa trong hệ thống"]
    rows_abbr = [
        ["API", "Application Programming Interface", "Giao diện lập trình ứng dụng"],
        ["DAG", "Directed Acyclic Graph", "Đồ thị có hướng không chu trình điều phối tác vụ"],
        ["dbt", "data build tool", "Công cụ chuyển đổi và mô hình hóa dữ liệu theo chuẩn kỹ thuật phần mềm"],
        ["EDW", "Enterprise Data Warehouse", "Kho dữ liệu doanh nghiệp"],
        ["IQR", "Interquartile Range", "Khoảng tứ phân vị trong phân tích thống kê"],
        ["Linear SVM", "Linear Support Vector Machine", "Máy vectơ hỗ trợ tuyến tính phân loại cảm xúc"],
        ["LR", "Logistic Regression", "Hồi quy Logistic làm mô hình cơ sở so sánh"],
        ["NLP", "Natural Language Processing", "Xử lý ngôn ngữ tự nhiên"],
        ["OLAP", "Online Analytical Processing", "Xử lý phân tích trực tuyến phục vụ báo cáo"],
        ["RDD", "Resilient Distributed Dataset", "Tập dữ liệu phân tán có khả năng phục hồi của Spark"],
        ["SPC", "Statistical Process Control", "Kiểm soát quá trình bằng thống kê"],
        ["TF-IDF", "Term Frequency - Inverse Document Frequency", "Tần suất từ kết hợp nghịch đảo tần suất văn bản"],
        ["TOC", "Table of Contents", "Mục lục tài liệu"],
        ["UI", "User Interface", "Giao diện người dùng tương tác"],
        ["UCSD", "University of California San Diego", "Đại học California tại San Diego, đơn vị công bố tập dữ liệu"],
    ]
    add_styled_table(doc, headers_abbr, rows_abbr, col_widths=[3.0, 6.0, 7.0])

    doc.add_page_break()
